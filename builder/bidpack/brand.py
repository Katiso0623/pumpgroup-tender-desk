"""Pump Group SA branding for every generated page.

PDF pages: the full-page letterhead (logo, "We don't sell pumps, we sell solutions",
blue accents, company footer) is drawn behind the content.
Word pages: built on the company's own Letter of Authority template, so the
letterhead sits in the header exactly as in their documents. The MD's signature
image in that template is always removed: signatures are done by hand.
"""
from __future__ import annotations

import base64
import copy
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Twips
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader

HERE = Path(__file__).resolve().parent.parent
TEMPLATES = HERE / "templates"

# Brand colours sampled from the letterhead
CHARCOAL = HexColor("#3D4045")
BLUE = HexColor("#1AA3DD")
INK = HexColor("#1E2326")
GREY = HexColor("#5E6670")
RULE = HexColor("#C9D1D8")
DOCX_BLUE = RGBColor(0x1A, 0xA3, 0xDD)
DOCX_CHARCOAL = RGBColor(0x3D, 0x40, 0x45)

# Area of an A4 page that the letterhead artwork leaves free (fractions of width/height from the top-left)
SAFE = {"left": 0.09, "right": 0.91, "top": 0.175, "bottom": 0.80}


def _file(name: str) -> Path:
    p = TEMPLATES / name
    if p.exists():
        return p
    b64 = TEMPLATES / (name + ".b64")
    if b64.exists():
        p.write_bytes(base64.b64decode(b64.read_text()))
        return p
    raise FileNotFoundError(f"Missing template {name}: copy the Project's bidpack/templates files")


def letterhead_image() -> Path:
    return _file("pump_group_letterhead.jpg")


def draw_letterhead(c, W: float = A4[0], H: float = A4[1]):
    """Draw the full-page letterhead on a reportlab canvas (portrait A4)."""
    c.drawImage(ImageReader(str(letterhead_image())), 0, 0, W, H)


def safe_box(W: float = A4[0], H: float = A4[1]):
    """(x0, y_top, x1, y_bottom) in reportlab coordinates."""
    return W * SAFE["left"], H * (1 - SAFE["top"]), W * SAFE["right"], H * (1 - SAFE["bottom"])


# ------------------------------------------------------------------- Word
def new_document(top_in: float = 1.85, bottom_in: float = 2.2) -> Document:
    """A blank Word document on the Pump Group letterhead (template body cleared, signature image removed)."""
    doc = Document(str(_file("letter_template.docx")))
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)
    for s in doc.sections:
        s.top_margin = Twips(int(top_in * 1440))
        s.bottom_margin = Twips(int(bottom_in * 1440))
        s.left_margin = Twips(1300)
        s.right_margin = Twips(1300)
        s.different_first_page_header_footer = False
    # the template has an empty "first page" header: make the letterhead header apply to every page
    sect = doc.sections[0]
    sect.first_page_header.is_linked_to_previous = True
    return doc


def para(doc, text: str = "", size: float = 10.5, bold: bool = False, align=None, color=None, space_after: float = 6,
         italic: bool = False, caps: bool = False):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text.upper() if caps else text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(space_after)
    if bold and size >= 11 and len(text) < 120:
        p.paragraph_format.keep_with_next = True
    return p


def runs(p, parts: list[tuple[str, bool]], size: float = 10.5):
    """parts: [(text, bold)]"""
    for text, bold in parts:
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
    return p


def shade(cell, hex_fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(doc, header: list[str], rows: list[list[str]], widths_in: list[float] | None = None, size: float = 9.5,
          header_fill: str = "3D4045"):
    from docx.shared import Inches
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, header_fill)
    from docx.oxml import OxmlElement as _OE
    from docx.oxml.ns import qn as _qn
    for row in rows:
        tr = t.add_row()
        trPr = tr._tr.get_or_add_trPr()
        cs = _OE("w:cantSplit")
        trPr.append(cs)
        cells = tr.cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v if v is not None else ""))
            r.font.size = Pt(size)
    if widths_in:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        t.autofit = False
        grid = t._tbl.tblGrid
        cols = grid.findall(qn("w:gridCol"))
        for i, w in enumerate(widths_in):
            if i < len(cols):
                cols[i].set(qn("w:w"), str(int(w * 1440)))
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for row in t.rows:
            for i, w in enumerate(widths_in):
                row.cells[i].width = Inches(w)
    return t


def signature_block(doc, lines: list[str], label: str = "Signature"):
    para(doc, "", space_after=18)
    para(doc, "_______________________________", space_after=0)
    for i, l in enumerate(lines):
        para(doc, l, bold=(i == 0), space_after=0)
    para(doc, "", space_after=4)


def save_docx_and_pdf(doc, out_docx: str) -> str | None:
    """Save the Word file and a PDF next to it (LibreOffice). Returns the PDF path or None."""
    import subprocess
    import sys
    doc.save(out_docx)
    outdir = str(Path(out_docx).parent)
    helper = Path("/mnt/skills/public/docx/scripts/office/soffice.py")
    cmd = ([sys.executable, str(helper)] if helper.exists() else ["soffice"]) + ["--headless", "--convert-to", "pdf", "--outdir", outdir, out_docx]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=240)
    except Exception:
        return None
    pdf = Path(out_docx).with_suffix(".pdf")
    return str(pdf) if pdf.exists() else None
