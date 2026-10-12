"""Fill the bidder details in a Word (.docx) tender document.

Only blank answer spaces are written to: an empty table cell next to a recognised
label, or a run of underscores / dots after a recognised label. Nothing that is
already written is overwritten. Signatures, dates, witnesses and declarations are
never filled. Every action is reported so it can be checked against the blank copy.
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field, asdict

from docx import Document
from docx.shared import Pt, RGBColor

from .common import field_value, is_blank, match_field

LEADER = r"(?:_{3,}|\.{4,}|…{2,}|(?:\.\s){4,}\.?)"
INLINE_RE = re.compile(r"(?P<label>[A-Za-z][A-Za-z0-9 ()/&'\-,]{2,90}?)\s*[:\-]?\s*(?P<blank>" + LEADER + r")")


@dataclass
class FillReport:
    filled: list = field(default_factory=list)      # what we wrote
    missing: list = field(default_factory=list)     # recognised label, but the profile has no value
    unknown: list = field(default_factory=list)     # blank spaces we did not recognise
    directors_rows: int = 0

    def as_dict(self):
        return asdict(self)


def _unique_cells(row):
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(c)
    return out


def _first_run_font(par):
    for r in par.runs:
        if r.text.strip():
            return r.font
    return None


def _write_cell(cell, value, like_cell=None, color=None):
    par = cell.paragraphs[0]
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    for extra in cell.paragraphs[1:]:
        if is_blank(extra.text):
            extra._p.getparent().remove(extra._p)
    run = par.add_run(value)
    src = None
    if like_cell is not None and like_cell.paragraphs:
        src = _first_run_font(like_cell.paragraphs[0])
    if src is not None:
        run.font.name = src.name
        run.font.size = src.size or Pt(10)
    else:
        run.font.size = Pt(10)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _replace_span(par, start, end, value, color=None):
    """Replace characters [start, end) of the paragraph text with value, run-aware."""
    pos = 0
    inserted = False
    for run in list(par.runs):
        t = run.text
        a, b = pos, pos + len(t)
        pos = b
        if b <= start or a >= end:
            continue
        cut_from = max(start, a) - a
        cut_to = min(end, b) - a
        if not inserted:
            # split the run: printed text keeps its look, only the inserted value is coloured
            import copy as _copy
            before, after = t[:cut_from], t[cut_to:]
            run.text = before
            new_r = _copy.deepcopy(run._r)
            run._r.addnext(new_r)
            from docx.text.run import Run as _Run
            val = _Run(new_r, run._parent)
            val.text = " " + value + " "
            val.font.underline = False
            if color:
                val.font.color.rgb = RGBColor.from_string(color)
            if after:
                tail_r = _copy.deepcopy(run._r)
                new_r.addnext(tail_r)
                _Run(tail_r, run._parent).text = after
            inserted = True
        else:
            run.text = t[:cut_from] + t[cut_to:]
    return inserted


def _fill_paragraph(par, profile, report, where, color):
    text = par.text
    if not text or not re.search(LEADER, text):
        return
    matches = list(INLINE_RE.finditer(text))
    for m in reversed(matches):  # right to left keeps earlier offsets valid
        label = m.group("label").strip()
        # use only the last clause of the label ("... 1.2 Name of bidder")
        label_tail = re.split(r"[.;:]\s+|\d+(?:\.\d+)+\s+", label)[-1].strip()
        key = match_field(label_tail) or match_field(label)
        if not key:
            if not re.search(r"sign|date|witness|initial", label, re.I):
                context = re.sub(r"\s+", " ", text[:m.start("blank")]).strip()[-90:]
                report.unknown.append({"where": where, "label": context or label_tail[:90]})
            continue
        value = field_value(profile, key)
        if not value:
            report.missing.append({"where": where, "label": label_tail[:90], "key": key})
            continue
        if _replace_span(par, m.start("blank"), m.end("blank"), value, color):
            report.filled.append({"where": where, "label": label_tail[:90], "key": key, "value": value})


DIRECTOR_HEADER = re.compile(r"full name|name of (director|member|partner|shareholder)", re.I)
ID_HEADER = re.compile(r"identity|id number|id no", re.I)


def _parse_directors(text):
    out = []
    for line in (text or "").splitlines():
        parts = [p.strip() for p in re.split(r"[,;|\t]", line) if p.strip()]
        if not parts:
            continue
        name = parts[0]
        idno = next((p for p in parts[1:] if re.fullmatch(r"\d{13}", p.replace(" ", ""))), "")
        share = next((p for p in parts[1:] if "%" in p), "")
        out.append({"name": name, "id": idno, "share": share})
    return out


def _fill_directors(table, profile, report, where, color):
    rows = table.rows
    if len(rows) < 2:
        return
    header = [c.text.strip() for c in _unique_cells(rows[0])]
    if not any(DIRECTOR_HEADER.search(h) for h in header) or not any(ID_HEADER.search(h) for h in header):
        return
    directors = _parse_directors(profile.get("directors", ""))
    if not directors:
        report.missing.append({"where": where, "label": "Directors / members table", "key": "directors"})
        return
    name_col = next(i for i, h in enumerate(header) if DIRECTOR_HEADER.search(h))
    id_col = next(i for i, h in enumerate(header) if ID_HEADER.search(h))
    share_col = next((i for i, h in enumerate(header) if re.search(r"%|share|interest", h, re.I)), None)
    di = 0
    for r in rows[1:]:
        cells = _unique_cells(r)
        if di >= len(directors):
            break
        if len(cells) <= max(name_col, id_col):
            continue
        if not is_blank(cells[name_col].text):
            continue
        d = directors[di]
        _write_cell(cells[name_col], d["name"], color=color)
        if d["id"] and is_blank(cells[id_col].text):
            _write_cell(cells[id_col], d["id"], color=color)
        if share_col is not None and d["share"] and share_col < len(cells) and is_blank(cells[share_col].text):
            _write_cell(cells[share_col], d["share"], color=color)
        di += 1
        report.directors_rows += 1
    if di:
        report.filled.append({"where": where, "label": "Directors / members table", "key": "directors", "value": f"{di} row(s)"})


def _fill_table(table, profile, report, where, color):
    _fill_directors(table, profile, report, where, color)
    for ri, row in enumerate(table.rows):
        cells = _unique_cells(row)
        for ci, cell in enumerate(cells):
            label = cell.text.strip()
            if not label or is_blank(label):
                continue
            # a label cell followed by an empty answer cell
            if ci + 1 < len(cells) and is_blank(cells[ci + 1].text) and len(label) <= 140:
                key = match_field(label)
                if key:
                    value = field_value(profile, key)
                    w = f"{where}, row {ri + 1}"
                    if value:
                        _write_cell(cells[ci + 1], value, like_cell=cell, color=color)
                        report.filled.append({"where": w, "label": label[:90], "key": key, "value": value})
                    else:
                        report.missing.append({"where": w, "label": label[:90], "key": key})
                    continue
                if len(label) <= 60 and not re.search(r"sign|date|witness|initial|stamp|yes|no\b", label, re.I) and ci == 0:
                    report.unknown.append({"where": f"{where}, row {ri + 1}", "label": label[:90]})
            for pi, par in enumerate(cell.paragraphs):
                _fill_paragraph(par, profile, report, f"{where}, row {ri + 1}", color)
            for ti, inner in enumerate(cell.tables):
                _fill_table(inner, profile, report, f"{where}.{ti + 1}", color)


def fill_docx(src: str, dst: str, profile: dict, color: str | None = "C00000") -> FillReport:
    """Fill bidder details into a copy of src, saved as dst, in red (default) so the copy is a filling guide:
    the user prints the original and writes the same entries in pen."""
    doc = Document(src)
    report = FillReport()
    heading = "start of document"
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph
            par = Paragraph(block, doc)
            t = par.text.strip()
            if t and len(t) < 90 and (t.isupper() or re.match(r"^(sbd|mbd|t2\.\d|c\d\.\d|schedule|form|part|section|annex)", t, re.I)):
                heading = t[:80]
            _fill_paragraph(par, profile, report, heading, color)
        elif tag == "tbl":
            from docx.table import Table
            _fill_table(Table(block, doc), profile, report, f"table under '{heading}'", color)
    for s in doc.sections:
        for part in (s.header, s.footer):
            for par in part.paragraphs:
                _fill_paragraph(par, profile, report, "header/footer", color)
    doc.save(dst)
    return report
